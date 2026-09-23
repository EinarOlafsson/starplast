"""Build the Starplast slide viewer, editable PowerPoint, PDF and GitHub pages.

One content source drives native text/shapes in both PDF and PowerPoint. Poppler
renders the exact PDF to slide images and thumbnails, so the repository, viewer
and downloadable deck agree. Requires reportlab, python-pptx, Pillow and pdftoppm.
The viewer design follows spaCR's MIT-licensed readme_deck_viewer.html.
"""
from pathlib import Path
import json
import math
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).parent))
from deck_content import SLIDES

OUT = ROOT / 'docs/deck'
W, H = 1600, 900
BG, PANEL, FG, MUTED, ACCENT, EDGE = '#0e1116', '#191f27', '#f4f6f8', '#a8b0ba', '#2ec4b6', '#353e48'


class Drawing:
    """Draw consistent native elements in a PDF and an editable PowerPoint."""
    def __init__(self):
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from pptx import Presentation
        from pptx.util import Pt
        self.metrics = pdfmetrics
        for name, filename in [('regular','Regular'),('bold','Bold'),('light','Light')]:
            pdfmetrics.registerFont(TTFont(name, str(OUT / f'fonts/OpenSans-{filename}.ttf')))
        self.pdf = canvas.Canvas(str(OUT / 'starplast_deck.pdf'), pagesize=(W,H))
        self.pdf.setAuthor('Einar Olafsson')
        self.pdf.setTitle('Starplast: introduction and practical guide')
        self.ppt = Presentation()
        self.ppt.slide_width, self.ppt.slide_height = Pt(W*.6), Pt(H*.6)
        self.ppt.core_properties.author = 'Einar Olafsson'
        self.ppt.core_properties.last_modified_by = 'Einar Olafsson'
        self.ppt.core_properties.title = 'Starplast: introduction and practical guide'
        self.number = 0

    @staticmethod
    def pt(value):
        from pptx.util import Pt
        return Pt(value*.6)

    @staticmethod
    def rgb(value):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(value.lstrip('#'))

    def shape(self, x,y,w,h,fill=PANEL,stroke=EDGE,rounding=False,ellipse=False):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml.xmlchemy import OxmlElement
        self.pdf.setFillColor(fill)
        self.pdf.setStrokeColor(stroke or fill)
        self.pdf.setLineWidth(1.6)
        if ellipse: self.pdf.ellipse(x,H-y-h,x+w,H-y,fill=1,stroke=bool(stroke))
        elif rounding: self.pdf.roundRect(x,H-y-h,w,h,18,fill=1,stroke=bool(stroke))
        else: self.pdf.rect(x,H-y-h,w,h,fill=1,stroke=bool(stroke))
        kind = MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.ROUNDED_RECTANGLE if rounding else MSO_SHAPE.RECTANGLE
        s=self.page.shapes.add_shape(kind,*map(self.pt,(x,y,w,h)))
        s.fill.solid();s.fill.fore_color.rgb=self.rgb(fill)
        if stroke: s.line.color.rgb=self.rgb(stroke);s.line.width=self.pt(1.6)
        else: s.line.fill.background()
        if rounding: s.adjustments[0]=.08
        s._element.spPr.append(OxmlElement('a:effectLst'))

    def line(self,x1,y1,x2,y2,color=EDGE,width=2):
        from pptx.enum.shapes import MSO_CONNECTOR
        self.pdf.setStrokeColor(color);self.pdf.setLineWidth(width)
        self.pdf.line(x1,H-y1,x2,H-y2)
        s=self.page.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*map(self.pt,(x1,y1,x2,y2)))
        s.line.color.rgb=self.rgb(color);s.line.width=self.pt(width)

    def text(self,x,y,w,content,size=30,color=FG,weight='regular',limit=None,link=None):
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        font='Courier' if weight=='mono' else weight
        lines=[]
        for paragraph in str(content).split('\n'):
            words=paragraph.split();current=''
            for word in words:
                proposed=(current+' '+word).strip()
                if self.metrics.stringWidth(proposed,font,size)>w and current:
                    lines.append(current);current=word
                else: current=proposed
            lines.append(current)
        leading=size*1.32
        height=len(lines)*leading
        if limit and height>limit:
            raise ValueError(f'Slide {self.number}: text exceeds {limit}px: {content!r} ({height:.0f}px)')
        if y+height>826:
            raise ValueError(f'Slide {self.number}: text enters footer: {content!r}')
        self.pdf.setFillColor(color);self.pdf.setFont(font,size)
        for i,line in enumerate(lines):self.pdf.drawString(x,H-y-size-i*leading,line)
        box=self.page.shapes.add_textbox(*map(self.pt,(x,y,w,height+5)))
        tf=box.text_frame;tf.clear();tf.word_wrap=False
        # LibreOffice shrinks spAutoFit text boxes around their centre, moving
        # left-aligned text. Fixed bounds preserve the authored coordinates.
        tf.auto_size=MSO_AUTO_SIZE.NONE
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment=PP_ALIGN.LEFT
            p.text=line;p.font.name='Consolas' if weight=='mono' else 'Open Sans'
            p.font.size=self.pt(size);p.font.bold=weight=='bold';p.font.color.rgb=self.rgb(color)
            p.line_spacing=self.pt(leading);p.space_after=0;p.space_before=0
            if link:
                for r in p.runs:r.hyperlink.address=link
        if link:self.pdf.linkURL(link,(x,H-y-height,x+w,H-y),relative=0,thickness=0)
        return height

    def image(self,name,x,y,w,h):
        from PIL import Image
        path=OUT/'assets'/name
        with Image.open(path) as im:iw,ih=im.size
        scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale
        dx,dy=x+(w-dw)/2,y+(h-dh)/2
        self.pdf.drawImage(str(path),dx,H-dy-dh,dw,dh,mask='auto')
        self.page.shapes.add_picture(str(path),*map(self.pt,(dx,dy,dw,dh)))

    def start(self, spec):
        self.number+=1;self.page=self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.shape(0,0,W,H,fill=BG,stroke=None)
        if spec['layout']!='cover':
            self.text(70,36,1460,spec['section'],22,ACCENT,'bold')
            self.text(70,82,1460,spec['title'],53,weight='bold',limit=145)
            if spec.get('lead') and spec['layout']!='end':
                self.text(70,165,1460,spec['lead'],29,MUTED,limit=80)
        self.page.notes_slide.notes_text_frame.text=plain(spec)+'\n\nSource: '+spec['source']

    def finish(self,spec):
        from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
        if spec.get('note') and spec['layout']!='cover':
            self.text(70,752,1460,spec['note'],24,MUTED,limit=66)
        self.line(70,834,1530,834)
        # Footer is intentionally below the content safety boundary.
        self.pdf.setFont('regular',18);self.pdf.setFillColor(MUTED)
        for x,content in [(70,'STARPLAST  /  0.43 GUIDE'),(1335,f'{self.number:02d} / {len(SLIDES):02d}')]:
            self.pdf.drawString(x,H-872,content)
            box=self.page.shapes.add_textbox(*map(self.pt,(x,850,350,30)))
            box.text_frame.auto_size=MSO_AUTO_SIZE.NONE
            p=box.text_frame.paragraphs[0];p.text=content;p.font.name='Open Sans'
            p.alignment=PP_ALIGN.LEFT
            p.font.size=self.pt(18);p.font.color.rgb=self.rgb(MUTED)
        self.pdf.linkURL(spec['source'],(70,18,800,62),relative=0,thickness=0)
        self.pdf.showPage()

    def save(self):
        self.pdf.save();self.ppt.save(OUT/'starplast_deck.pptx')


def plain(spec, include_title=True):
    """Readable transcript shared by slide pages, PowerPoint notes and the viewer."""
    lines=([spec['title']] if include_title else [])+[spec.get('lead','')]
    if spec.get('code'):lines.append(spec['code'])
    for title,body in spec.get('cards',[]):lines.append(title+': '+body)
    if spec.get('rows'):
        lines.append(' | '.join(spec['columns']))
        lines.extend(' | '.join(row) for row in spec['rows'])
    lines.append(spec.get('note',''))
    return '\n\n'.join(line for line in lines if line)


def side_cards(d,cards,x=1050,y=250,w=470,gap=156):
    for i,(title,body) in enumerate(cards):
        top=y+i*gap
        d.text(x,top,w,title,30,ACCENT,'bold',limit=80)
        d.text(x,top+47,w,body,26,FG,limit=103)


def render_slide(d,s):
    """Lay out one idea using a small set of deliberately distinct compositions."""
    d.start(s);kind=s['layout'];cards=s.get('cards',[])
    if kind=='cover':
        d.text(80,170,900,'Starplast',112,weight='bold')
        d.text(85,350,830,s['lead'],43,ACCENT,limit=190)
        d.image('logo.png',1080,175,390,390)
        d.text(85,680,1400,s['note'],29,MUTED)
        d.text(85,732,1400,'Einar Olafsson  ·  github.com/EinarOlafsson/starplast',24,MUTED)
    elif kind in ('cards','split'):
        n=len(cards);gap=30;w=(1460-gap*(n-1))/n
        for i,(title,body) in enumerate(cards):
            x=70+i*(w+gap)
            d.shape(x,260,w,450,rounding=True)
            d.line(x+30,298,x+95,298,ACCENT,3)
            ht=d.text(x+30,330,w-60,title,36,weight='bold',limit=150)
            d.text(x+30,355+ht,w-60,body,30,MUTED,limit=295-ht)
    elif kind=='stats':
        for i,(title,body) in enumerate(cards):
            x=70+i*500
            d.shape(x,265,460,430,rounding=True)
            size=92
            while d.metrics.stringWidth(title,'bold',size)>400:size-=2
            d.text(x+30,326,400,title,size,ACCENT,'bold',limit=145)
            d.text(x+30,490,400,body,34,limit=150)
    elif kind=='flow':
        n=len(cards);w=(1460-28*(n-1))/n
        for i,(title,body) in enumerate(cards):
            x=70+i*(w+28)
            d.shape(x,265,w,435,rounding=True)
            d.text(x+26,290,w-52,f'{i+1:02d}',66,ACCENT,'light')
            ht=d.text(x+26,395,w-52,title,31,weight='bold',limit=93)
            d.text(x+26,415+ht,w-52,body,27,MUTED,limit=250-ht)
            if i<n-1:d.line(x+w+3,477,x+w+25,477,ACCENT,2)
    elif kind=='image':
        side_cards(d,cards,x=70,y=255,w=465,gap=153)
        d.shape(580,245,950,470,fill='#11161c',rounding=True)
        d.image(s['image'],595,260,920,440)
    elif kind=='fullimage':
        d.image(s['image'],70,230,1460,490)
    elif kind=='code':
        d.shape(70,252,960,463,rounding=True)
        size=25 if len(s['code'].splitlines())>5 else 31
        d.text(100,285,900,s['code'],size,FG,'mono',limit=407)
        side_cards(d,cards,x=1080,y=263,w=440,gap=150 if len(cards)==3 else 205)
    elif kind=='table':
        xs=[70,455,1010];ws=[350,515,500]
        for x,w,title in zip(xs,ws,s['columns']):d.text(x+15,252,w-20,title,26,ACCENT,'bold')
        for i,row in enumerate(s['rows']):
            top=310+i*79
            d.shape(70,top-8,1460,76,fill=PANEL if i%2==0 else BG,stroke=None)
            for x,w,value in zip(xs,ws,row):d.text(x+15,top,w-24,value,25,limit=68)
    elif kind=='network':
        points=[(310,310),(175,445),(330,485),(480,350),(550,530),(355,650),(130,630)]
        for a,b in [(0,1),(0,2),(0,3),(1,2),(1,6),(2,4),(2,5),(3,4),(4,5),(5,6)]:
            d.line(*points[a],*points[b],ACCENT if a==2 or b==2 else MUTED,2)
        for i,(x,y) in enumerate(points):d.shape(x-13,y-13,26,26,fill=ACCENT if i==2 else FG,stroke=None,ellipse=True)
        d.text(70,701,620,'Illustrative network; line meaning comes from its layer.',22,MUTED)
        side_cards(d,cards,x=780,y=270,w=730,gap=151)
    elif kind=='selection':
        for i,(title,body) in enumerate(cards):
            x=70+i*750
            d.shape(x,255,710,445,rounding=True)
            d.text(x+25,280,660,title,33,ACCENT,'bold')
            for j in range(27):
                px=x+65+(j*67%565);py=385+(j*47%155)
                d.shape(px,py,10,10,fill=ACCENT if 170<px-x<450 else MUTED,stroke=None,ellipse=True)
            if i==0:
                polygon=[(x+180,365),(x+445,385),(x+470,530),(x+230,562),(x+160,475),(x+180,365)]
                for a,b in zip(polygon,polygon[1:]):d.line(*a,*b,FG,2)
            else:
                # Unfilled outlines keep the point cloud visible.
                for angle in range(0,360,5):
                    a,b=math.radians(angle),math.radians(angle+5)
                    d.line(x+325+115*math.cos(a),460+85*math.sin(a),x+325+115*math.cos(b),460+85*math.sin(b),FG,2)
            d.text(x+25,586,660,body,26,MUTED,limit=85)
    elif kind=='folds':
        for i,(title,body) in enumerate(cards):
            x=70+i*500
            d.text(x,262,460,title,40,ACCENT,'bold')
            for row in range(3):
                for col in range(4):
                    d.shape(x+20+col*90,345+row*66,32,32,fill=[FG,MUTED,ACCENT][i],stroke=None,ellipse=True)
            d.text(x,588,440,body,29,limit=130)
        d.text(70,690,1400,'Circles represent genes. Related genes stay within one group.',23,MUTED)
    elif kind=='coverage':
        sequence=s['metric']=='sequence';n=8064 if sequence else 1210
        d.text(70,285,880,f'{n:,} / 8,140',86,ACCENT,'bold')
        d.text(70,415,870,'Toxoplasma genes with this feature view',32,MUTED)
        d.shape(70,505,860,48,fill=EDGE,stroke=None)
        d.shape(70,505,860*n/8140,48,fill=ACCENT,stroke=None)
        d.text(70,595,850,f'{100*n/8140:.1f}% coverage in the bundled release',34)
        side_cards(d,cards)
    elif kind=='metrics':
        metric_chart(d,s['metric'])
        side_cards(d,cards,x=1060,w=465,y=258,gap=153)
    elif kind=='themes':
        for i,name in enumerate(('dark','light')):
            x=70+i*750
            d.image('map-'+name+'.png',x,245,710,445)
            d.text(x,700,710,name.capitalize()+' theme',27,MUTED)
    elif kind=='end':
        d.text(70,180,1400,s['lead'],45,ACCENT,'mono')
        for i,(title,url) in enumerate(cards):
            x=70+(i%3)*500;y=310+(i//3)*192
            d.shape(x,y,460,155,rounding=True)
            d.text(x+25,y+25,410,title,33,weight='bold',link=url,limit=92)
            d.text(x+25,y+96,410,'Open resource  >',24,ACCENT,link=url)
    else:raise ValueError(kind)
    d.finish(s)


def metric_chart(d,kind):
    """Draw directly from the committed benchmark artifacts, with labelled units."""
    import csv
    results=ROOT/'results/inference_043'
    with (results/'summary.csv').open() as f:rows=list(csv.DictReader(f))
    if kind=='calibration':
        with (results/'reliability_bins.csv').open() as f:bins=list(csv.DictReader(f))
        data=[r for r in bins if r['method']=='boosted' and r['features']=='published_af3_esm'
              and float(r['bin_lower'])>=.7]
        d.text(70,243,850,'Mean confidence / observed accuracy',26,MUTED)
        for i,r in enumerate(data):
            y=310+i*123
            d.text(70,y,245,f"{float(r['bin_lower']):.1f}-{float(r['bin_lower'])+.1:.1f}",29)
            for j,(field,color) in enumerate([('mean_confidence',MUTED),('accuracy',ACCENT)]):
                value=float(r[field]);d.shape(330,y+j*37,510*value,26,fill=color,stroke=None)
                d.text(870,y+j*37-4,130,f'{value:.3f}',25,color)
        return
    target='compartment' if kind=='localization' else 'protein_ibaq_log2'
    field='macro_f1' if kind=='localization' else 'rmse'
    entries=[('prior','published','Prior'),('linear','published','Linear')]
    if kind=='localization':entries+=[('umap','published','UMAP 20D')]
    entries += [('boosted','published','Boosted'),('boosted','published_af3_esm','Boosted + AF3 + ESM')]
    maximum=.5 if kind=='localization' else 4
    d.text(70,243,880,'Macro F1 (higher is better)' if kind=='localization' else 'RMSE in log2 units (lower is better)',26,MUTED)
    for i,(method,view,label) in enumerate(entries):
        record=next(r for r in rows if r['task']==target and r['method']==method and r['features']==view)
        value=float(record[field]);y=315+i*(73 if kind=='localization' else 90)
        d.text(70,y,290,label,25,limit=70)
        color=ACCENT if view.endswith('esm') else MUTED
        d.shape(380,y+5,500*value/maximum,30,fill=color,stroke=None)
        d.text(905,y,130,f'{value:.3f}',27,color)


def publish():
    """Render portable images and emit the viewer, transcript and linked GitHub pages."""
    assert 1<=len(SLIDES)<=50
    OUT.mkdir(parents=True,exist_ok=True)
    drawing=Drawing()
    for spec in SLIDES:render_slide(drawing,spec)
    drawing.save()
    for folder,width,quality in [('slides',1600,90),('thumbs',320,85)]:
        destination=OUT/folder;destination.mkdir(exist_ok=True)
        subprocess.run(['pdftoppm','-jpeg','-jpegopt',f'quality={quality},optimize=y',
                        '-scale-to-x',str(width),'-scale-to-y','-1',
                        str(OUT/'starplast_deck.pdf'),str(destination/'slide')],check=True)
        for p in destination.glob('slide-*.jpg'):
            p.replace(destination/f'slide_{int(p.stem.split("-")[-1]):02d}.jpg')
    deck=dict(count=len(SLIDES),aspect=16/9,title='Starplast: introduction and practical guide',
              author='Einar Olafsson',guide_version='0.43',slides=[])
    pages=OUT/'pages';pages.mkdir(exist_ok=True)
    transcript=['# Starplast: slide transcript','A guide to the 0.43 release.']
    for i,s in enumerate(SLIDES,1):
        transcript.extend([f'## {i:02d}. {s["title"]}',plain(s,False),f'[Supporting documentation]({s["source"]})'])
        row=dict(image=f'slides/slide_{i:02d}.jpg',thumb=f'thumbs/slide_{i:02d}.jpg',
                 title=s['title'],section=s['section'],text=plain(s,False),source=s['source'],animations=[])
        deck['slides'].append(row)
        previous=len(SLIDES) if i==1 else i-1;next_slide=1 if i==len(SLIDES) else i+1
        page=(f'[![{s["title"]}](../slides/slide_{i:02d}.jpg)](https://einarolafsson.github.io/starplast/deck/#{i})\n\n'
              f'[← Back]({previous:02d}.md) · {i} / {len(SLIDES)} · [Next →]({next_slide:02d}.md)\n\n'
              f'## {s["title"]}\n\n{plain(s,False)}\n\n[Supporting documentation]({s["source"]})\n')
        (pages/f'{i:02d}.md').write_text(page)
    (OUT/'slides.json').write_text(json.dumps(deck,indent=2)+'\n')
    (OUT/'transcript.md').write_text('\n\n'.join(transcript)+'\n')
    template=(ROOT/'scripts/deck_viewer.html').read_text()
    (OUT/'index.html').write_text(template.replace('__SLIDES__',json.dumps(deck).replace('</','<\\/')))
    print(f'Built {len(SLIDES)} slides: {OUT}')


if __name__=='__main__':publish()
